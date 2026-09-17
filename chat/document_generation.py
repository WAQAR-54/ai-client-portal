"""Turns one assistant message's content into a downloadable Word/Excel/
PowerPoint/PDF file - gated by the "document_generation" Plan feature flag
(see governance/models.py::KNOWN_FEATURE_FLAGS, enforced in
chat/views.py::export_message_document).

Distinct from chat/export.py, which exports a whole CONVERSATION's
transcript (every message, plain sender/timestamp formatting) - this
generates one nicely-structured document from a SINGLE message's content,
parsed from the same sanitized HTML chat/markdown_utils.render_markdown()
already produces for on-screen rendering. Walking that already-bleach-
cleaned HTML (via lxml, already a dependency) means this never has to
parse the model's raw Markdown a second, differently-behaved way, and
never touches unsanitized model output.
"""

from io import BytesIO

from lxml import html as lxml_html
from xhtml2pdf import pisa

from chat.markdown_utils import render_markdown

_HEADING_TAGS = {"h1", "h2", "h3", "h4", "h5", "h6"}


def _parsed_blocks(message_content):
    """Yields (tag, value) for each top-level block in the message's
    rendered HTML: ("h1".."h6", text), ("li", text), ("table", rows) where
    rows is a list of cell-text lists, or ("p", text) for anything else."""
    html = render_markdown(message_content)
    if not html.strip():
        return
    root = lxml_html.fragment_fromstring(f"<div>{html}</div>")
    for el in root.iterchildren():
        tag = el.tag
        if tag in _HEADING_TAGS:
            text = el.text_content().strip()
            if text:
                yield (tag, text)
        elif tag in ("ul", "ol"):
            for li in el.findall("li"):
                text = li.text_content().strip()
                if text:
                    yield ("li", text)
        elif tag == "table":
            rows = [[cell.text_content().strip() for cell in tr.findall("./*")] for tr in el.findall(".//tr")]
            rows = [r for r in rows if r]
            if rows:
                yield ("table", rows)
        else:
            text = el.text_content().strip()
            if text:
                yield ("p", text)


def extract_document_title(content, fallback="Document"):
    """The artifact panel's title - the reply's own first heading (the
    model is prompted, via chat/prompts.py::DOCUMENT_OUTPUT_HINT, to always
    lead with one), reusing this module's own markdown parsing rather than
    a separate regex pass over raw text. Falls back to `fallback` (the
    provisional title set from the user's prompt, see chat/views.py::
    post_message) when the reply has no heading at all - never raises."""
    for tag, value in _parsed_blocks(content):
        if tag in _HEADING_TAGS:
            return value[:200]
    return fallback


def render_message_docx(message) -> bytes:
    from docx import Document

    document = Document()
    for tag, value in _parsed_blocks(message.content):
        if tag in _HEADING_TAGS:
            document.add_heading(value, level=min(int(tag[1]), 4))
        elif tag == "li":
            document.add_paragraph(value, style="List Bullet")
        elif tag == "table":
            width = max(len(row) for row in value)
            table = document.add_table(rows=len(value), cols=width)
            table.style = "Light Grid Accent 1"
            for r, row in enumerate(value):
                for c, cell_text in enumerate(row):
                    table.cell(r, c).text = cell_text
        else:
            document.add_paragraph(value)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_message_xlsx(message) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Document"
    row_idx = 1
    for tag, value in _parsed_blocks(message.content):
        if tag == "table":
            for row in value:
                for c, cell_text in enumerate(row, start=1):
                    sheet.cell(row=row_idx, column=c, value=cell_text)
                row_idx += 1
            row_idx += 1  # blank separator row after an embedded table
        else:
            text = f"• {value}" if tag == "li" else value
            cell = sheet.cell(row=row_idx, column=1, value=text)
            if tag in _HEADING_TAGS:
                cell.font = Font(bold=True, size=max(11, 16 - int(tag[1])))
            row_idx += 1
    for column_cells in sheet.columns:
        values = [str(c.value) for c in column_cells if c.value is not None]
        length = max((len(v) for v in values), default=10)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 60)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_message_pptx(message) -> bytes:
    """One slide per heading (its following paragraphs/bullets as body
    text), one slide per table. Content with no heading at all still gets
    a single slide titled "Generated document" rather than being dropped."""
    from pptx import Presentation

    presentation = Presentation()
    content_layout = presentation.slide_layouts[1]  # "Title and Content"

    def _new_slide(title):
        slide = presentation.slides.add_slide(content_layout)
        slide.shapes.title.text = title
        return slide

    def _set_body(slide, lines):
        if not lines:
            return
        body = slide.placeholders[1].text_frame
        body.clear()
        body.paragraphs[0].text = lines[0]
        for line in lines[1:]:
            body.add_paragraph().text = line

    slide = None
    body_lines = []
    for tag, value in _parsed_blocks(message.content):
        if tag in _HEADING_TAGS:
            _set_body(slide, body_lines)
            slide = _new_slide(value)
            body_lines = []
        elif tag == "table":
            _set_body(slide, body_lines)
            slide = _new_slide("Table")
            body_lines = [" | ".join(row) for row in value]
            _set_body(slide, body_lines)
            slide, body_lines = None, []
        else:
            if slide is None:
                slide = _new_slide("Generated document")
            body_lines.append(value)
    _set_body(slide, body_lines)

    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def render_message_pdf(message) -> bytes:
    html = (
        "<html><head><style>"
        "body { font-family: Helvetica, sans-serif; font-size: 11pt; }"
        "h1, h2, h3, h4 { color: #222; }"
        "table { border-collapse: collapse; width: 100%; }"
        "td, th { border: 1px solid #ccc; padding: 4px 8px; }"
        "</style></head><body>" + render_markdown(message.content) + "</body></html>"
    )
    buffer = BytesIO()
    status = pisa.CreatePDF(html, dest=buffer)
    if status.err:
        raise ValueError(f"PDF generation failed for message {message.id}")
    return buffer.getvalue()
